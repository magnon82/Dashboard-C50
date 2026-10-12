"""
Importación única de recetarios.

Lee una vez el master, los Excel de cocina y barra, y las fichas 2024.
Escribe data/recetario/snapshot.json. El dashboard no vuelve a abrir Drive.

Después de esto, el único dato que cambia es el costo del insumo, y lo
actualiza scripts/update_recetario_costos_facturas.py con el precio de factura.

Uso (desde la raíz del repo):
  python scripts/import_recetario.py
"""

from __future__ import annotations

import io
import json
import re
import sys
import unicodedata
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
from google.oauth2.credentials import Credentials
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "recetario" / "snapshot.json"
TOKEN = ROOT / "ingestor" / "token.json"

MASTER_ID = "1WK2bPjGxz46PS6WE36LwOfhZ_9CznEQht4ityFG9YiY"
FICHA_ID = "1YtoV1ltB_2HoB4oUDk096HoDUaOqoKBXEbVh6dEsvSI"
COCINA_XLSX = "1KWltzzbDC-yeFix0rITOijp4OKDkhzNm"
BARRA_XLSX = "1uRCAYf5YRuDtYjTsJwGuJBTenm-vwOpW"

# Nombres del JSON de Eventos que no coinciden con el PDF vigente.
DISPLAY = {
    "ensalada natural": "Queso fundido natural",
    "champinones": "Queso fundido con champiñones",
    "chistorra": "Queso fundido con chistorra",
}

# Menú vigente → nombre en el recetario de julio.
ALIASES = {
    "flores del huerto": ["flor de huitlacoche", "flor de hutlacoche"],
    "queso fundido natural": ["queso fundido natural", "queso fundido"],
    "queso fundido con champinones": [
        "queso fundido con champinones",
        "queso fundido champis",
        "queso fundido champinones",
    ],
    "queso fundido con chistorra": [
        "queso fundido con chistorra",
        "queso fundido chistorra",
    ],
    "coliflor rostizada": ["coliflor asada", "coliflor rostizada"],
    "tartara de cecina": ["carne tartara de cecina", "tartara de cecina"],
    "patatas carranza": ["papas cambray bravas", "patatas carranza"],
    "ensalada de tomates cherry": ["ensalada de jitomates heirloom"],
    "hamburguesa don venustiano": ["hamburguesa de diezmillo", "hamburguesa de diezmil"],
    "nuestro chemita": ["filete chemita", "nuestro chemita"],
    "costillar nuxi": ["costillar nuxi"],
    "pork belly en pipian verde": ["pork belly al pipian", "pork belly en pipian"],
    "manchamanteles": ["pollo en mancha manteles", "manchamanteles"],
    "helado de dulce de leche": ["helado de temporada", "helado de dulce de leche"],
    "ravioles de requeson y hongos": ["ravioles de requeson", "ravioles con requeson"],
    "emparedado de pulled pork": ["emparedado pulled pork"],
    "bunuelos con helado de tepache": ["bunuelos"],
    "pan de elote": ["pan de elote con helado de cardamomo", "pan de elote"],
    "cheesecake de la casa": ["cheescake", "cheesecake"],
    "mousse de 3 chocolates": ["mouse de 3 chocolates", "mousse de 3 chocolates"],
    "esquites de santa rosa": ["esquites serranos"],
    "sopa de hongos al carbon": ["sopa de hongo al carbon"],
    "ensalada de tomates cherry": ["ensalada de jitomates herloom", "ensalada de jitomates heirloom"],
    "camarones a las rocas": ["camaron a las rocas"],
    "espresso": ["cafe espresso"],
    "margarita": ["margarita de limon"],
    "mezcalita": ["mezcalita de limon"],
    "bloody mary": ["blooddy mery", "blody mery"],
    "baby carranza": ["baby mango"],
    "ensalada cesar al carbon": ["ensalada cesar al carbon"],
    "aperol spritz": ["apperol spritz", "aperol"],
    "clericot jarra": ["jarra de clericot"],
    "clericot copa": ["copa de clericot"],
    "cafe americano": ["cafe americano", "americano"],
    "espresso descafeinado": ["cafe americano descafeinado", "americano descafeinado"],
}


def norm(value: object) -> str:
    s = str(value or "").replace("ñ", "n").replace("Ñ", "n")
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().lower()
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def slug(value: str) -> str:
    return norm(value).replace(" ", "-")[:80] or "x"


def parse_num(value: object) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    s = str(value).strip().replace("$", "").replace(" ", "")
    if not s or s.upper() in {"NA", "N/A", "-"}:
        return None
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def download(drive, file_id: str) -> io.BytesIO:
    req = drive.files().get_media(fileId=file_id, supportsAllDrives=True)
    buf = io.BytesIO()
    dl = MediaIoBaseDownload(buf, req)
    done = False
    while not done:
        _, done = dl.next_chunk()
    buf.seek(0)
    return buf


def sheet_matrix(ws, max_row: int = 90, max_col: int = 14) -> list[list]:
    rows: list[list] = []
    for row in ws.iter_rows(max_row=max_row, max_col=max_col, values_only=True):
        rows.append(list(row))
    return rows


def find_sheet(wb, *names: str):
    wanted = {norm(n) for n in names}
    for name in wb.sheetnames:
        if norm(name) in wanted:
            return wb[name]
    return None


def header_map(row: list) -> dict[str, int]:
    found: dict[str, int] = {}
    for i, cell in enumerate(row):
        key = norm(cell)
        if key:
            found[key] = i
    return found


def col(headers: dict[str, int], *keys: str) -> int | None:
    for key in keys:
        if key in headers:
            return headers[key]
    for key in keys:
        for h, i in headers.items():
            if key in h:
                return i
    return None


def cell(row: list, idx: int | None):
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def parse_precios(ws) -> dict[int, dict]:
    rows = sheet_matrix(ws, max_row=1200, max_col=12)
    header_at = None
    headers: dict[str, int] = {}
    for i, row in enumerate(rows):
        h = header_map(row)
        if "producto" in h or any(k.startswith("producto") for k in h):
            header_at = i
            headers = h
            break
    out: dict[int, dict] = {}
    if header_at is None:
        return out
    i_name = col(headers, "producto")
    i_unit = col(headers, "unidad de costeo", "unidad de medida")
    i_sup = col(headers, "1er proveedor", "1er. proveedor")
    for offset, row in enumerate(rows[header_at + 1 :], start=header_at + 2):
        name = cell(row, i_name)
        if not name or not str(name).strip():
            continue
        out[offset] = {
            "name": str(name).strip(),
            "nameKey": norm(name),
            "unit": str(cell(row, i_unit) or "").strip(),
            "supplier": str(cell(row, i_sup) or "").strip(),
        }
    return out


def formula_ref(value: object):
    if not isinstance(value, str) or not value.startswith("="):
        return None
    m = re.search(r"PRECIOS!\$?([A-Z]+)\$?(\d+)", value, re.I)
    if m:
        return ("precios", int(m.group(2)))
    m = re.search(r"'?([^'!]+)'?!\$?[A-Z]+\$?(\d+)", value)
    if m:
        return ("sheet", m.group(1).strip())
    return None


def row_refs(row: list):
    precio = None
    sheet = None
    for value in row:
        ref = formula_ref(value)
        if not ref:
            continue
        if ref[0] == "precios" and precio is None:
            precio = ref[1]
        elif ref[0] == "sheet" and sheet is None:
            sheet = ref[1]
    return precio, sheet


def text_after(rows: list[list], label: str) -> str:
    lab = norm(label)
    for row in rows:
        for i, value in enumerate(row):
            raw = str(value or "").strip()
            if not raw or lab not in norm(raw):
                continue
            if ":" in raw:
                rest = raw.split(":", 1)[1].strip()
                if rest:
                    return rest
            for nxt in row[i + 1 :]:
                piece = str(nxt or "").strip()
                if piece and not piece.startswith("="):
                    return piece
    return ""


def yield_info(rows: list[list]) -> tuple[float, str, float]:
    for row in rows:
        for i, value in enumerate(row):
            if "rendimiento" not in norm(value):
                continue
            unit = ""
            nums: list[float] = []
            for nxt in row[i + 1 :]:
                if isinstance(nxt, str) and nxt.startswith("="):
                    continue
                num = parse_num(nxt)
                if num is not None and not (isinstance(nxt, str) and norm(nxt).isalpha()):
                    # skip text that parsed empty; numbers count
                    if not isinstance(nxt, str) or re.search(r"\d", nxt):
                        nums.append(num)
                        continue
                piece = str(nxt or "").strip()
                if piece and not unit and not piece.startswith("="):
                    unit = piece
            qty = nums[0] if nums else 1.0
            factor = nums[1] if len(nums) > 1 else 1.0
            if qty <= 0:
                qty = 1.0
            if factor <= 0:
                factor = 1.0
            return qty, unit, factor
    return 1.0, "porcion", 1.0


def parse_cocina_recipe(code: str, rows: list[list]) -> dict | None:
    blob = norm(" ".join(str(c or "") for row in rows[:8] for c in row[:4]))
    if "nombre de la receta" not in blob and "sub receta" not in blob:
        return None
    name = text_after(rows, "nombre de la receta")
    if not name:
        return None
    kind = "subreceta" if "sub receta" in blob else "platillo"
    header_idx = None
    for i, row in enumerate(rows):
        keys = header_map(row)
        has_ing = any(k.startswith("ingrediente") for k in keys)
        has_qty = any(k == "cantidad" or k.startswith("cantidad") for k in keys)
        if has_ing and has_qty:
            header_idx = i
            break
    if header_idx is None:
        return None
    headers = header_map(rows[header_idx])
    i_name = col(headers, "ingredientes", "ingrediente")
    i_qty = col(headers, "cantidad")
    i_unit = col(headers, "unidad")
    i_sub = col(headers, "sub")
    lines = []
    for row in rows[header_idx + 1 :]:
        iname = cell(row, i_name)
        qty = parse_num(cell(row, i_qty))
        if not iname or not str(iname).strip() or str(iname).startswith("="):
            continue
        if qty is None or qty == 0:
            continue
        label = str(iname).strip()
        if norm(label) in {"ingredientes", "gramaje total"}:
            continue
        precio_row, sheet_ref = row_refs(row)
        sub = cell(row, i_sub)
        sub_code = None
        if isinstance(sub, (int, float)) and not isinstance(sub, bool):
            sub_code = str(int(sub))
        elif isinstance(sub, str) and norm(sub).replace(" ", "").isdigit():
            sub_code = str(int(float(sub)))
        if sheet_ref and norm(sheet_ref) not in {"precios", "indice", "menu"}:
            sub_code = sub_code or sheet_ref
        lines.append(
            {
                "name": label,
                "qty": qty,
                "unit": str(cell(row, i_unit) or "").strip(),
                "precioRow": precio_row,
                "subCode": sub_code,
            }
        )
    yqty, yunit, factor = yield_info(rows)
    return {
        "code": code,
        "name": name,
        "nameKey": norm(name),
        "kind": kind,
        "category": text_after(rows, "categoria"),
        "yieldQty": yqty,
        "yieldUnit": yunit or "porcion",
        "orderFactor": factor,
        "method": text_after(rows, "metodo")[:2000],
        "lines": lines,
    }


def parse_barra_recipe(code: str, rows: list[list]) -> dict | None:
    blob = norm(" ".join(str(c or "") for row in rows[:6] for c in row[:6]))
    if "titulo de la receta" not in blob:
        return None
    name = text_after(rows, "titulo de la receta")
    if not name:
        return None
    header_idx = None
    for i, row in enumerate(rows):
        keys = header_map(row)
        if any(k.startswith("producto") for k in keys) and any("cantidad" in k for k in keys):
            header_idx = i
            break
    if header_idx is None:
        return None
    headers = header_map(rows[header_idx])
    i_name = col(headers, "producto")
    i_qty = col(headers, "cantidad")
    i_unit = col(headers, "unidad de receta", "unidad")
    lines = []
    for row in rows[header_idx + 1 :]:
        iname = cell(row, i_name)
        qty = parse_num(cell(row, i_qty))
        if not iname or not str(iname).strip() or str(iname).startswith("="):
            continue
        if qty is None or qty == 0:
            continue
        unit = str(cell(row, i_unit) or "").strip()
        # El Excel dice ML, pero 0.15 es litros (el precio de costeo es por litro).
        if norm(unit) in {"ml", "millilitro", "mililitro"} and qty < 10:
            unit = "lt"
        precio_row, sheet_ref = row_refs(row)
        sub_code = None
        if sheet_ref and norm(sheet_ref) not in {
            "precios",
            "indice",
            "menu",
            "vinos",
            "destilados",
            "cocteles",
            "cocteles nuevos",
        }:
            sub_code = sheet_ref
        lines.append(
            {
                "name": str(iname).strip(),
                "qty": qty,
                "unit": unit,
                "precioRow": precio_row,
                "subCode": sub_code,
            }
        )
    yqty, yunit, factor = yield_info(rows)
    if not yunit:
        yunit = "porcion"
        yqty = factor or 1
    return {
        "code": code,
        "name": name,
        "nameKey": norm(name),
        "kind": "coctel",
        "category": text_after(rows, "linea") or "Coctelería",
        "yieldQty": yqty or 1,
        "yieldUnit": yunit,
        "orderFactor": factor or 1,
        "method": "",
        "lines": lines,
    }


def load_recipes(wb, area: str) -> list[dict]:
    skip = {"menu", "indice", "precios", "machote", "machote receta", "destilados", "vinos"}
    recipes = []
    parser = parse_cocina_recipe if area == "cocina" else parse_barra_recipe
    for name in wb.sheetnames:
        if norm(name) in skip or norm(name).startswith("cervezas"):
            continue
        rec = parser(name, sheet_matrix(wb[name]))
        if rec and rec["lines"]:
            rec["area"] = area
            rec["id"] = f"{area}:{name}"
            recipes.append(rec)
    return recipes


def master_rows(sheets, title: str, a1: str) -> list[list]:
    res = (
        sheets.spreadsheets()
        .values()
        .get(spreadsheetId=MASTER_ID, range=f"'{title}'!{a1}")
        .execute()
    )
    return res.get("values") or []


def money_pair(row: list, start: int) -> dict | None:
    def at(i):
        return row[i] if i < len(row) else ""

    supplier = str(at(start) or "").strip()
    if not supplier or norm(supplier) in {"na", "n a"}:
        return None
    return {
        "supplier": supplier,
        "costingUnit": norm(at(start + 4)) or norm(at(start + 1)),
        "costingPrice": parse_num(at(start + 5)),
    }


def insumos_from_master(sheets) -> tuple[list[dict], list[dict]]:
    suppliers: dict[str, dict] = {}
    indice = master_rows(sheets, "Indice", "A1:N400")
    for row in indice[1:]:
        name = str(row[0] or "").strip() if row else ""
        if not name:
            continue
        key = norm(name)
        suppliers[key] = {
            "id": f"prov:{slug(name)}",
            "name": name,
            "nameKey": key,
            "category": str(row[1] or "").strip() if len(row) > 1 else "",
            "razonSocial": str(row[2] or "").strip() if len(row) > 2 else "",
            "rfc": str(row[3] or "").strip() if len(row) > 3 else "",
            "credito": str(row[11] or "").strip() if len(row) > 11 else "",
            "periodo": str(row[12] or "").strip() if len(row) > 12 else "",
        }

    insumos: list[dict] = []

    def add(area: str, name: str, category: str, offers: list[dict]):
        clean = [o for o in offers if o and o.get("costingPrice")]
        if not name or not clean:
            return
        first = clean[0]
        unit = first.get("costingUnit") or "pza"
        if unit in {"ml", "millilitro"}:
            unit = "lt"
        insumos.append(
            {
                "id": f"{area}:{slug(name)}",
                "area": area,
                "name": name.strip(),
                "nameKey": norm(name),
                "category": category.strip(),
                "supplier": first.get("supplier") or "",
                "costingUnit": unit or "pza",
                "costingPrice": round(float(first["costingPrice"]), 4),
                "priceSource": "master",
                "priceAsOf": None,
                "facturaEmisor": None,
                "offers": [
                    {
                        "rank": i + 1,
                        "supplier": o.get("supplier") or "",
                        "costingUnit": (o.get("costingUnit") or unit or "pza"),
                        "costingPrice": round(float(o["costingPrice"]), 4),
                    }
                    for i, o in enumerate(clean[:3])
                ],
            }
        )

    alimentos = master_rows(sheets, "Alimentos", "A1:X1200")
    for row in alimentos[1:]:
        name = str(row[0] or "").strip() if row else ""
        category = str(row[1] or "").strip() if len(row) > 1 else ""
        offers = []
        # 1er proveedor empieza en col C (index 2): proveedor, unidad, presentacion, precio, unidad costeo, precio costeo
        for start in (2, 8, 16):
            if start >= len(row):
                continue
            supplier = str(row[start] or "").strip()
            if not supplier:
                continue
            offers.append(
                {
                    "supplier": supplier,
                    "costingUnit": norm(row[start + 4]) if start + 4 < len(row) else "",
                    "costingPrice": parse_num(row[start + 5]) if start + 5 < len(row) else None,
                }
            )
        add("cocina", name, category, offers)

    barra = master_rows(sheets, "Barra", "A1:H1200")
    for row in barra[1:]:
        if len(row) < 8:
            continue
        name = str(row[1] or "").strip()
        category = str(row[2] or "").strip()
        price = parse_num(row[7])
        unit = norm(row[4]) if len(row) > 4 else "lt"
        if unit in {"ml", "millilitro"}:
            unit = "lt"
        add(
            "barra",
            name,
            category,
            [{"supplier": str(row[3] or "").strip(), "costingUnit": unit or "lt", "costingPrice": price}],
        )
    return list(suppliers.values()), insumos


def index_insumos(insumos: list[dict]) -> dict[str, list[dict]]:
    by: dict[str, list[dict]] = {}
    for item in insumos:
        by.setdefault(item["nameKey"], []).append(item)
    return by


def match_insumo(name: str, area: str, by_key: dict[str, list[dict]]) -> dict | None:
    key = norm(name)
    if not key:
        return None
    pool = by_key.get(key) or []
    same = [i for i in pool if i["area"] == area]
    if same:
        return same[0]
    if pool:
        return pool[0]
    cands = []
    for items in by_key.values():
        for item in items:
            other = item["nameKey"]
            if len(key) < 6 or len(other) < 6:
                continue
            if key in other or other in key:
                cands.append(item)
    same = [i for i in cands if i["area"] == area]
    use = same or cands
    keys = {i["id"] for i in use}
    if len(keys) == 1:
        return use[0]
    return None


def match_ranked(name: str, area: str, by_key: dict[str, list[dict]]) -> dict | None:
    """Exacto primero. El nombre escrito en la receta gana sobre la fila de PRECIOS."""
    key = norm(name)
    pool = by_key.get(key) or []
    same = [i for i in pool if i["area"] == area]
    if same:
        return same[0]
    if pool:
        return pool[0]
    return None


def link_lines(recipes: list[dict], precios: dict[int, dict], insumos_by: dict[str, list[dict]]):
    by_code = {r["code"]: r for r in recipes}
    for rec in recipes:
        for line in rec["lines"]:
            catalog = ""
            if line.get("precioRow") and line["precioRow"] in precios:
                catalog = precios[line["precioRow"]]["name"]
            found = match_ranked(line["name"], rec["area"], insumos_by)
            if not found and catalog:
                found = match_ranked(catalog, rec["area"], insumos_by)
            if not found:
                found = match_insumo(line["name"], rec["area"], insumos_by)
            if not found and catalog:
                found = match_insumo(catalog, rec["area"], insumos_by)
            line["insumoId"] = found["id"] if found else None
            sub = line.get("subCode")
            line["subrecipeId"] = None
            if sub and sub in by_code and by_code[sub]["id"] != rec["id"]:
                line["subrecipeId"] = by_code[sub]["id"]
            line.pop("precioRow", None)
            line.pop("subCode", None)


def load_menu() -> list[dict]:
    alimentos = json.loads((ROOT / "docs" / "eventos-menus" / "alimentos-c50-items.json").read_text(encoding="utf-8"))
    bebidas = json.loads((ROOT / "docs" / "eventos-menus" / "bebidas-c50-items.json").read_text(encoding="utf-8"))
    items = []
    for raw in alimentos:
        name = raw["name"]
        shown = DISPLAY.get(norm(name), name)
        cat = "Platillo"
        desc = str(raw.get("description") or "")
        if "·" in desc:
            cat = desc.split("·")[-1].strip().capitalize()
        items.append(
            {
                "sku": raw["sku"],
                "area": "cocina",
                "name": shown,
                "nameKey": norm(shown),
                "category": cat,
                "sellPrice": float(raw["unit_price"]),
            }
        )
    for raw in bebidas:
        items.append(
            {
                "sku": raw["sku"],
                "area": "barra",
                "name": raw["name"],
                "nameKey": norm(raw["name"]),
                "category": bebida_category(raw["sku"], raw["name"]),
                "sellPrice": float(raw["unit_price"]),
            }
        )
    return items


def bebida_category(sku: str, name: str) -> str:
    s = sku.upper()
    n = norm(name)
    if s.startswith("VIN") or s.startswith("BEB-PROS"):
        return "Vino"
    if s.startswith("TEQ"):
        return "Tequila"
    if s.startswith("RON"):
        return "Ron"
    if s.startswith("WHI") or s.startswith("WHIS"):
        return "Whisky"
    if s.startswith("GIN"):
        return "Ginebra"
    if s.startswith("MEZ"):
        return "Mezcal"
    if s.startswith("BRA") or s.startswith("BRAND"):
        return "Brandy"
    if s.startswith("COG"):
        return "Cognac"
    if s.startswith("VOD"):
        return "Vodka"
    if s.startswith("LIC"):
        return "Licor"
    if "ART" in s:
        return "Cerveza artesanal"
    if "CERV" in s or "TAR" in s or "CLA" == s.split("-")[-1] and "CERV" in s:
        return "Cerveza"
    if any(x in s for x in ("JUG", "LIM", "NAR", "PINADA", "CONGA", "AGUA", "COCA", "GING", "MUND", "SPR", "FRES", "BOOST", "PERR")):
        return "Sin alcohol"
    if n in {"americano", "espresso", "espresso descafeinado", "espresso doble", "capuchino", "latte", "te", "cafe americano"} or "cafe" in n or n in {"capuchino", "latte"}:
        return "Café"
    return "Coctelería"


PORCION_CATS = {
    "Vino",
    "Tequila",
    "Ron",
    "Whisky",
    "Ginebra",
    "Mezcal",
    "Brandy",
    "Cognac",
    "Vodka",
    "Licor",
    "Cerveza",
    "Cerveza artesanal",
    "Sin alcohol",
}


MATCH_STOP = {"copeo", "botella", "copa", "jarra", "ml", "pieza", "nacional"}


def stem_word(word: str) -> str:
    if len(word) > 4 and word.endswith("es"):
        return word[:-2]
    if len(word) > 4 and word.endswith("s"):
        return word[:-1]
    return word


def match_words(name: str) -> list[str]:
    return [
        stem_word(w)
        for w in norm(name).split()
        if w not in MATCH_STOP and len(w) > 1
    ]


def words_cover(needle: str, hay: str) -> bool:
    nw = match_words(needle)
    hw = set(match_words(hay))
    if not nw:
        return False
    if len(nw) == 1 and len(nw[0]) < 6:
        return False
    return all(w in hw for w in nw)


def find_recipe(item: dict, recipes: list[dict]) -> dict | None:
    keys = [item["nameKey"], *ALIASES.get(item["nameKey"], [])]
    bare = " ".join(match_words(item["name"]))
    if bare and bare not in keys:
        keys.append(bare)
    by = {}
    for rec in recipes:
        if rec["area"] != item["area"] or rec["kind"] == "subreceta":
            continue
        by.setdefault(rec["nameKey"], rec)
    for key in keys:
        if key in by:
            return by[key]
        if key in ALIASES:
            for alt in ALIASES[key]:
                if alt in by:
                    return by[alt]
    cands = []
    for rec in recipes:
        if rec["area"] != item["area"] or rec["kind"] == "subreceta":
            continue
        if words_cover(item["name"], rec["name"]) or words_cover(rec["name"], item["name"]):
            cands.append(rec)
    ids = {c["id"] for c in cands}
    if len(ids) == 1:
        return cands[0]
    return None


def portion_for(item: dict, insumos_by: dict[str, list[dict]]) -> dict | None:
    if item["category"] not in PORCION_CATS:
        return None
    sku = item["sku"].upper()
    bottle = sku.endswith("-B")
    found = match_insumo(" ".join(match_words(item["name"])), "barra", insumos_by)
    if not found:
        tokens = match_words(item["name"])
        cands = []
        for items in insumos_by.values():
            for ins in items:
                if ins["area"] != "barra":
                    continue
                words = set(ins["nameKey"].split())
                if tokens and all(t in words for t in tokens):
                    cands.append(ins)
        uniq = {c["id"]: c for c in cands}
        if len(uniq) == 1:
            found = next(iter(uniq.values()))
        elif len(uniq) > 1:
            found = sorted(uniq.values(), key=lambda i: len(i["nameKey"]))[0]
    if not found:
        return {"insumoId": None, "qty": None, "unit": None, "label": None}
    if item["category"] == "Vino":
        qty, label = (0.75, "Botella 750 ml") if bottle else (0.15, "Copeo 150 ml")
    elif item["category"] == "Licor":
        qty, label = (0.06, "Copeo 60 ml") if not bottle else (0.7, "Botella")
    elif item["category"] in {"Cerveza", "Cerveza artesanal"}:
        ml = 473 if item["category"] == "Cerveza artesanal" and "hercules" in item["nameKey"] else 355
        qty, label = ml / 1000, f"Pieza {ml} ml"
    elif item["category"] == "Sin alcohol":
        qty, label = 0.355, "Pieza 355 ml"
    else:
        if bottle:
            qty, label = 0.7, "Botella"
        else:
            qty, label = 0.045, "Copeo 45 ml"
    return {"insumoId": found["id"], "qty": qty, "unit": "lt", "label": label}


def ficha_value(rows: list[list], label: str) -> str:
    return text_after(rows, label)


def load_fichas(sheets) -> dict[str, dict]:
    meta = (
        sheets.spreadsheets()
        .get(spreadsheetId=FICHA_ID, fields="sheets(properties(title))")
        .execute()
    )
    titles = [
        (sh.get("properties") or {}).get("title")
        for sh in meta.get("sheets") or []
    ]
    titles = [t for t in titles if t and norm(t) != "indice"]
    ranges = [f"'{t}'!A1:L30" for t in titles]
    out: dict[str, dict] = {}
    for i in range(0, len(ranges), 20):
        chunk = ranges[i : i + 20]
        res = (
            sheets.spreadsheets()
            .values()
            .batchGet(spreadsheetId=FICHA_ID, ranges=chunk)
            .execute()
        )
        for block in res.get("valueRanges") or []:
            rows = block.get("values") or []
            name = ficha_value(rows, "nombre")
            if not name:
                continue
            out[norm(name)] = {
                "version": ficha_value(rows, "version"),
                "descripcion": ficha_value(rows, "descripcion"),
                "categoria": ficha_value(rows, "categoria"),
                "observaciones": ficha_value(rows, "observaciones"),
                "tipoCocina": ficha_value(rows, "tipo de cocina"),
                "tiempoCiclo": ficha_value(rows, "tiempo de ciclo"),
                "alergenos": ficha_value(rows, "alergenos"),
                "presentacion": ficha_value(rows, "presentacion"),
                "equipo": ficha_value(rows, "equipo a utilizar"),
                "acompanamientos": ficha_value(rows, "acompanamientos"),
                "procesos": ficha_value(rows, "procesos a destacar"),
                "historia": ficha_value(rows, "historia"),
            }
    return out


def blank_ficha(name: str, category: str) -> dict:
    return {
        "version": "",
        "descripcion": "",
        "categoria": category,
        "observaciones": "",
        "tipoCocina": "",
        "tiempoCiclo": "",
        "alergenos": "",
        "presentacion": "",
        "equipo": "",
        "acompanamientos": "",
        "procesos": "",
        "historia": "",
        "nombre": name,
    }


def attach_menu(menu, recipes, insumos_by, fichas):
    linked_ids = set()
    for item in menu:
        rec = find_recipe(item, recipes)
        item["recipeId"] = rec["id"] if rec else None
        item["portion"] = None
        if rec:
            linked_ids.add(rec["id"])
        elif item["area"] == "barra":
            item["portion"] = portion_for(item, insumos_by)
        ficha = fichas.get(item["nameKey"])
        if not ficha:
            for alt in ALIASES.get(item["nameKey"], []):
                if alt in fichas:
                    ficha = fichas[alt]
                    break
        item["ficha"] = ficha or blank_ficha(item["name"], item["category"])
        item["ficha"]["nombre"] = item["name"]
        if rec and rec.get("method") and not item["ficha"].get("procesos"):
            item["ficha"]["procesos"] = rec["method"]
    for rec in recipes:
        if rec["kind"] == "subreceta" or any(
            line.get("subrecipeId") == rec["id"] for other in recipes for line in other["lines"]
        ):
            if rec["id"] not in linked_ids and rec["kind"] != "platillo":
                rec["status"] = "interno"
            elif rec["id"] in linked_ids:
                rec["status"] = "vigente"
            elif rec["kind"] == "subreceta":
                rec["status"] = "interno"
            else:
                rec["status"] = "archivado" if rec["id"] not in linked_ids else "vigente"
        else:
            rec["status"] = "vigente" if rec["id"] in linked_ids else "archivado"
    # Subrecetas referenciadas quedan internas aunque también se hayan marcado arriba.
    referenced = {
        line["subrecipeId"]
        for rec in recipes
        for line in rec["lines"]
        if line.get("subrecipeId")
    }
    for rec in recipes:
        if rec["id"] in referenced and rec["id"] not in linked_ids:
            rec["status"] = "interno"
        elif rec["id"] in linked_ids:
            rec["status"] = "vigente"


def slim_recipe(rec: dict) -> dict:
    return {
        "id": rec["id"],
        "area": rec["area"],
        "code": rec["code"],
        "name": rec["name"],
        "nameKey": rec["nameKey"],
        "kind": rec["kind"],
        "status": rec["status"],
        "category": rec.get("category") or "",
        "yieldQty": rec["yieldQty"],
        "yieldUnit": rec["yieldUnit"],
        "orderFactor": rec["orderFactor"],
        "method": rec.get("method") or "",
        "lines": [
            {
                "name": line["name"],
                "qty": line["qty"],
                "unit": line["unit"],
                "insumoId": line.get("insumoId"),
                "subrecipeId": line.get("subrecipeId"),
            }
            for line in rec["lines"]
        ],
    }


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    creds = Credentials.from_authorized_user_file(
        str(TOKEN),
        [
            "https://www.googleapis.com/auth/spreadsheets.readonly",
            "https://www.googleapis.com/auth/drive.readonly",
        ],
    )
    sheets = build("sheets", "v4", credentials=creds)
    drive = build("drive", "v3", credentials=creds)

    print("Master…")
    suppliers, insumos = insumos_from_master(sheets)
    insumos_by = index_insumos(insumos)
    print(f"  proveedores {len(suppliers)} insumos {len(insumos)}")

    print("Fichas…")
    fichas = load_fichas(sheets)
    print(f"  fichas con texto {len(fichas)}")

    print("Cocina…")
    wb = openpyxl.load_workbook(download(drive, COCINA_XLSX), data_only=False, read_only=True)
    precios_c = parse_precios(find_sheet(wb, "PRECIOS"))
    recipes = load_recipes(wb, "cocina")
    wb.close()
    link_lines(recipes, precios_c, insumos_by)
    print(f"  recetas {len(recipes)}")

    print("Barra…")
    wb = openpyxl.load_workbook(download(drive, BARRA_XLSX), data_only=False, read_only=True)
    precios_b = parse_precios(find_sheet(wb, "PRECIOS"))
    barra = load_recipes(wb, "barra")
    wb.close()
    link_lines(barra, precios_b, insumos_by)
    recipes.extend(barra)
    print(f"  recetas {len(barra)}")

    menu = load_menu()
    attach_menu(menu, recipes, insumos_by, fichas)

    cocina_menu = [m for m in menu if m["area"] == "cocina"]
    barra_menu = [m for m in menu if m["area"] == "barra"]
    print(
        "Menú cocina enlazado",
        sum(1 for m in cocina_menu if m["recipeId"]),
        "/",
        len(cocina_menu),
    )
    print("Sin receta cocina:")
    for m in cocina_menu:
        if not m["recipeId"]:
            print("  -", m["name"])
    print(
        "Menú barra con receta",
        sum(1 for m in barra_menu if m["recipeId"]),
        "con porción",
        sum(1 for m in barra_menu if (m.get("portion") or {}).get("insumoId")),
        "/",
        len(barra_menu),
    )
    print("Barra sin receta ni insumo:")
    for m in barra_menu:
        portion = m.get("portion") or {}
        if not m["recipeId"] and not portion.get("insumoId"):
            print("  -", m["category"], m["name"])

    snapshot = {
        "importedAt": datetime.now(timezone.utc).isoformat(),
        "sourceNote": "Copia local. Cocina y barra no vuelven a leer Drive. El costo del insumo se actualiza con facturas.",
        "suppliers": suppliers,
        "insumos": insumos,
        "recipes": [slim_recipe(r) for r in recipes],
        "menu": [
            {
                "sku": m["sku"],
                "area": m["area"],
                "name": m["name"],
                "nameKey": m["nameKey"],
                "category": m["category"],
                "sellPrice": m["sellPrice"],
                "recipeId": m["recipeId"],
                "portion": m["portion"],
                "ficha": m["ficha"],
            }
            for m in menu
        ],
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(snapshot, ensure_ascii=False, indent=2), encoding="utf-8")
    print("Escrito", OUT, "bytes", OUT.stat().st_size)


if __name__ == "__main__":
    main()
